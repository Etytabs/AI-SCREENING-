"""
Document Extractor - Extract text from PDF and DOCX files
Supports multiple document formats for grant proposals
"""

import os
from typing import Dict, Any, Optional
from io import BytesIO

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class DocumentExtractor:
    """Extract text from various document formats"""
    
    def __init__(self):
        self.supported_formats = []
        if PDF_AVAILABLE:
            self.supported_formats.append('pdf')
        if DOCX_AVAILABLE:
            self.supported_formats.append('docx')
    
    def extract(self, file_path: str, file_content: Optional[bytes] = None) -> Dict[str, Any]:
        """
        Extract text from a document
        
        Args:
            file_path: Path to the file or filename with extension
            file_content: Optional file content as bytes
            
        Returns:
            Dict with extracted text and metadata
        """
        file_ext = self._get_extension(file_path)
        
        if file_ext == 'pdf' and PDF_AVAILABLE:
            return self._extract_pdf(file_path, file_content)
        elif file_ext == 'docx' and DOCX_AVAILABLE:
            return self._extract_docx(file_path, file_content)
        elif file_ext == 'txt':
            return self._extract_text(file_path, file_content)
        else:
            return {
                "success": False,
                "error": f"Unsupported format: {file_ext}. Supported: {self.supported_formats}",
                "text": "",
                "metadata": {}
            }
    
    def _get_extension(self, file_path: str) -> str:
        """Get file extension"""
        return os.path.splitext(file_path)[1].lower().lstrip('.')
    
    def _extract_pdf(self, file_path: str, file_content: Optional[bytes] = None) -> Dict[str, Any]:
        """Extract text from PDF"""
        try:
            if file_content:
                # Extract from bytes
                pdf_file = BytesIO(file_content)
                pdf_reader = PyPDF2.PdfReader(pdf_file)
            else:
                # Extract from file path
                with open(file_path, 'rb') as f:
                    pdf_reader = PyPDF2.PdfReader(f)
            
            text = ""
            num_pages = len(pdf_reader.pages)
            
            for page_num, page in enumerate(pdf_reader.pages):
                text += page.extract_text() + "\n\n"
            
            return {
                "success": True,
                "text": text.strip(),
                "metadata": {
                    "format": "pdf",
                    "num_pages": num_pages,
                    "character_count": len(text)
                }
            }
        except Exception as e:
            return {
                "success": False,
                "error": f"PDF extraction failed: {str(e)}",
                "text": "",
                "metadata": {}
            }
    
    def _extract_docx(self, file_path: str, file_content: Optional[bytes] = None) -> Dict[str, Any]:
        """Extract text from DOCX"""
        try:
            if file_content:
                # Extract from bytes
                doc_file = BytesIO(file_content)
                doc = Document(doc_file)
            else:
                # Extract from file path
                doc = Document(file_path)
            
            text = ""
            paragraph_count = 0
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
                paragraph_count += 1
            
            return {
                "success": True,
                "text": text.strip(),
                "metadata": {
                    "format": "docx",
                    "paragraph_count": paragraph_count,
                    "character_count": len(text)
                }
            }
        except Exception as e:
            return {
                "success": False,
                "error": f"DOCX extraction failed: {str(e)}",
                "text": "",
                "metadata": {}
            }
    
    def _extract_text(self, file_path: str, file_content: Optional[bytes] = None) -> Dict[str, Any]:
        """Extract text from plain text file"""
        try:
            if file_content:
                text = file_content.decode('utf-8')
            else:
                with open(file_path, 'r', encoding='utf-8') as f:
                    text = f.read()
            
            return {
                "success": True,
                "text": text.strip(),
                "metadata": {
                    "format": "txt",
                    "character_count": len(text)
                }
            }
        except Exception as e:
            return {
                "success": False,
                "error": f"Text extraction failed: {str(e)}",
                "text": "",
                "metadata": {}
            }
    
    def get_supported_formats(self) -> list:
        """Get list of supported file formats"""
        return self.supported_formats.copy()
