from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from apps.api.dependencies import get_service
from apps.api.main import app
from services.grant_workflow.documents import extract_bytes
from services.grant_workflow.duplication_archive import (
    ArchiveValidationError,
    DuplicateProjectError,
    DuplicationArchive,
)
from services.grant_workflow.jobs import InlineJobRunner
from services.grant_workflow.models import Application, ApplicationDocument, DataOrigin
from services.grant_workflow.service import GrantWorkflowService
from services.sources.registry import SourceAccessStatus

ADMIN = {"X-User-Role": "GRANT_ADMINISTRATOR", "X-User-Id": "admin-library"}
REVIEWER = {"X-User-Role": "REVIEWER", "X-User-Id": "reviewer-library"}
PROPOSAL = b"We will compare precision irrigation sensors across sixty farms.\n\fMethods: randomized field trial with soil moisture monitoring."


@pytest.fixture
def archive(tmp_path):
    return DuplicationArchive(tmp_path / "library.sqlite3")


def _import(archive, **kwargs):
    return archive.import_document(
        filename="proposal.txt", data=PROPOSAL, title="Irrigation trial",
        source_type="funded_project", **kwargs,
    )


def test_import_persists_text_pages_and_metadata_after_restart(archive):
    project = _import(archive, reference="FUND-2025-1", year=2025, organization="Research Council", uploaded_by="admin-1")
    reopened = DuplicationArchive(archive.path)
    assert reopened.status() == SourceAccessStatus.AVAILABLE
    assert reopened.list_projects() == [project]
    record, = reopened.historical_records()
    assert record.record_id == project.id
    assert record.source_type == "funded_project" and record.outcome == "funded"
    assert record.data_origin == DataOrigin.UPLOADED
    assert record.document.pages[1].page_number == 2
    assert "randomized field trial" in record.document.pages[1].text
    assert record.document_meta.filename == "proposal.txt"
    assert record.provenance["reference"] == "FUND-2025-1"
    assert record.provenance["uploaded_by"] == "admin-1"
    assert "text" not in project.model_dump()


def test_reimport_cannot_create_duplicate_content_with_new_filename(archive):
    first = _import(archive)
    with pytest.raises(DuplicateProjectError) as error:
        archive.import_document(
            filename="renamed.txt", data=PROPOSAL.upper().replace(b" ", b"  "),
            title="A different title", source_type="funded_project",
        )
    assert error.value.project.id == first.id
    assert len(archive.list_projects()) == 1


def test_funded_import_preserves_distinct_provenance_from_existing_submission(archive):
    archive.snapshot_submission(
        application_id="app-1", grant_call_id="call-1", title="Submitted irrigation trial", text=PROPOSAL.decode(),
    )
    funded = _import(archive)
    previous = archive.import_document(
        filename="proposal.txt", data=PROPOSAL, title="Previously submitted trial", source_type="historical_application",
    )
    assert funded.id != previous.id
    assert len(archive.list_projects()) == 3


def test_concurrent_reimports_only_create_one_record(archive):
    def attempt(_):
        try:
            return _import(archive).id
        except DuplicateProjectError as exc:
            return exc.project.id

    with ThreadPoolExecutor(max_workers=4) as workers:
        ids = list(workers.map(attempt, range(8)))
    assert len(set(ids)) == 1
    assert len(archive.list_projects()) == 1


def test_snapshot_upserts_only_same_application_and_excludes_self(archive):
    args = dict(grant_call_id="call-old", title="Old proposal", text=PROPOSAL.decode(), reference="APP-1")
    first = archive.snapshot_submission(application_id="app-1", **args)
    other = archive.snapshot_submission(application_id="app-2", **args)
    edited = archive.snapshot_submission(application_id="app-1", **{**args, "text": "An updated narrative about different trial sites."})
    assert first.id == edited.id and first.created_at == edited.created_at
    assert other.id != first.id
    assert len(archive.list_projects()) == 2
    record, = archive.historical_records(exclude_application_id="app-1")
    assert record.application_id == "app-2" and record.grant_call_id == "call-old"


def test_application_snapshot_preserves_original_sources_and_skips_demo(archive):
    application = Application(grant_call_id="call-1", application_reference="REF-1", title="Irrigation")
    doc = extract_bytes("proposal.txt", PROPOSAL, "doc-1")
    meta = ApplicationDocument(
        id="doc-1", application_id=application.id, filename="proposal.txt",
        document_type="proposal", file_hash=doc.file_sha256,
        uploaded_by="admin", extraction_status="success",
    )
    project = archive.snapshot_application(application, [(meta, doc)], "Institute")
    assert project.application_id == application.id
    record, = archive.historical_records()
    assert record.document.source_id == "doc-1"
    assert record.document_meta.application_id == application.id
    assert record.provenance["submitted_at"] == application.submitted_at.isoformat()
    assert record.provenance["documents"][0]["pages"][1]["page_number"] == 2
    synthetic = application.model_copy(update={"id": "synthetic", "data_origin": DataOrigin.SYNTHETIC})
    assert archive.snapshot_application(synthetic, [(meta, doc)]) is None
    assert len(archive.list_projects()) == 1


def test_multiple_document_snapshot_preserves_lines_without_inventing_pages(archive):
    application = Application(grant_call_id="call-1", application_reference="REF-1")
    documents = []
    for index, text in enumerate([
        b"Title: Irrigation research\nWe monitor soil moisture across sixty farms.",
        b"Title: Trial protocol\nMethods include daily randomized sampling in dry seasons.",
    ]):
        document_id = f"doc-{index}"
        filename = f"proposal-{index}.txt"
        doc = extract_bytes(filename, text, document_id)
        meta = ApplicationDocument(
            id=document_id, application_id=application.id, filename=filename,
            document_type="proposal", file_hash=doc.file_sha256,
            uploaded_by="admin", extraction_status="success",
        )
        documents.append((meta, doc))
    archive.snapshot_application(application, documents)
    record, = DuplicationArchive(archive.path).historical_records()
    assert "Title: Irrigation research\nWe monitor" in record.text
    assert "Title: Trial protocol\nMethods include" in record.text
    assert record.document is None and record.document_meta is None


@pytest.mark.parametrize("filename,data", [
    ("empty.txt", b""), ("space.txt", b" \n\t"), ("binary.txt", b"\xff\xfe\x00"),
    ("binary.txt", b"before\x00after"), ("broken.pdf", b"not a PDF"),
    ("broken.docx", b"not a DOCX"), ("notes.csv", b"some text"),
])
def test_invalid_files_are_not_archived(archive, filename, data):
    with pytest.raises(ArchiveValidationError):
        archive.import_document(filename=filename, data=data, title="Invalid", source_type="funded_project")
    assert archive.list_projects() == []


def test_valid_docx_and_pdf_imports(archive):
    import io

    import fitz
    from docx import Document

    word = Document()
    word.add_paragraph("A funded investigation of mobile clinics in remote settlements.")
    buffer = io.BytesIO()
    word.save(buffer)
    word_record = archive.import_document(
        filename="clinic.docx", data=buffer.getvalue(), title="Mobile clinics", source_type="funded_project",
    )
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "A prior proposal on coastal erosion monitoring using drone imagery.")
        pdf_record = archive.import_document(
            filename="coasts.pdf", data=pdf.tobytes(), title="Coast monitoring", source_type="historical_application",
        )
    records = {r.record_id: r for r in archive.historical_records()}
    assert "mobile clinics" in records[word_record.id].text
    assert "coastal erosion" in records[pdf_record.id].document.pages[0].text


@pytest.fixture
def archive_client(archive):
    service = GrantWorkflowService(jobs=InlineJobRunner(), providers=[], duplication_archive=archive)
    app.dependency_overrides[get_service] = lambda: service
    yield TestClient(app)
    app.dependency_overrides.pop(get_service, None)


def test_import_api_authorization_validation_and_conflict(archive_client):
    files = {"file": ("proposal.txt", PROPOSAL, "text/plain")}
    data = {"title": "Irrigation trial", "source_type": "funded_project", "year": "2025"}
    assert archive_client.post("/api/v1/duplication/projects", files=files, data=data).status_code == 401
    assert archive_client.post("/api/v1/duplication/projects", files=files, data=data, headers=REVIEWER).status_code == 403
    imported = archive_client.post("/api/v1/duplication/projects", files=files, data=data, headers=ADMIN)
    assert imported.status_code == 201
    body = imported.json()
    assert body["data_origin"] == "UPLOADED" and body["text_length"] > 0
    assert body["source_type"] == "funded_project" and body["year"] == 2025
    assert "text" not in body
    assert archive_client.get("/api/v1/duplication/projects").json() == [body]
    duplicate = archive_client.post("/api/v1/duplication/projects", files=files, data=data, headers=ADMIN)
    assert duplicate.status_code == 409 and body["id"] in duplicate.json()["detail"]
    invalid = archive_client.post("/api/v1/duplication/projects", files=files, data={**data, "source_type": "invented"}, headers=ADMIN)
    assert invalid.status_code == 422
    empty_title = archive_client.post("/api/v1/duplication/projects", files=files, data={**data, "title": "  "}, headers=ADMIN)
    assert empty_title.status_code == 422
    assert len(archive_client.get("/api/v1/duplication/projects").json()) == 1


def test_configured_database_path(monkeypatch, tmp_path):
    selected = tmp_path / "nested" / "comparison.sqlite3"
    monkeypatch.setenv("AI_SCREENING_DUPLICATION_DB", str(selected))
    archive = DuplicationArchive()
    assert archive.path == Path(selected).resolve()
    assert not selected.exists()
    assert archive.status() == SourceAccessStatus.AVAILABLE
    assert selected.exists()


def test_unavailable_archive_does_not_break_service_construction(tmp_path):
    blocked = tmp_path / "not-a-directory"
    blocked.write_text("a file blocks creation of the parent folder")
    archive = DuplicationArchive(blocked / "library.sqlite3")
    service = GrantWorkflowService(jobs=InlineJobRunner(), providers=[], duplication_archive=archive)
    assert service.duplication_archive is archive
    assert archive.status() == SourceAccessStatus.UNAVAILABLE
