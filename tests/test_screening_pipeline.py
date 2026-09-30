import fitz
from docx import Document

from services.screening.pipeline import screen_document


def test_screening_rejects_unsupported_documents(tmp_path):
    path = tmp_path / "proposal.odt"
    path.write_bytes(b"demo")
    result = screen_document(str(path), "TEST-001")
    assert result["status"] == "unsupported_type"
    assert result["extraction"]["extraction_status"] == "unsupported_type"


def test_screening_processes_pdf_document(tmp_path):
    path = tmp_path / "proposal.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Abstract\nMethodology\nBudget")
    document.save(path)
    document.close()

    result = screen_document(str(path), "TEST-PDF")

    assert result["status"] == "screened"
    assert result["extraction"]["extraction_status"] == "success"
    assert result["extraction"]["page_count"] == 1
    assert result["completeness"]["total"] == 6
    assert result["human_review_required"] is True


def test_screening_processes_docx_document(tmp_path):
    path = tmp_path / "proposal.docx"
    document = Document()
    document.add_paragraph("Abstract")
    document.add_paragraph("Methodology")
    document.add_paragraph("Budget")
    document.save(path)

    result = screen_document(str(path), "TEST-DOCX")

    assert result["status"] == "screened"
    assert result["extraction"]["extraction_status"] == "success"
    assert result["completeness"]["total"] == 6
    assert result["human_review_required"] is True
